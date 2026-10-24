"""Importa las planillas Excel de pautas de la nutricionista como plantillas de pauta.

Lee la carpeta de planificación (Omnívoros/, Vegetarianos/...) y convierte cada hoja
de menú en una PlantillaPauta: porciones por grupo, distribución por tiempo de comida
y menú en formato de planilla ({tiempo: {opciones: [{plato, lineas: [{porcion, texto}]}]}}).

Formatos soportados:
  - "Horizontal": columna B con las porciones ("1 Cereal") y por cada opción una
    columna con el plato y otra con el alimento ("2 Reb pan molde Castaños").
  - Vertical: hoja PORCIONES (porciones y alimentos por opción) + hoja MENÚ (platos).

Uso:
  python scripts/importar_planillas.py revisar <carpeta> <salida.json>
  python scripts/importar_planillas.py cargar  <carpeta> <usuario>     (usa DATABASE_URL)
"""
import glob
import json
import os
import re
import sys
import unicodedata
from fractions import Fraction
from typing import Dict, List, Optional, Tuple

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from routers.pautas import _GRUPOS_MACROS as GRUPOS_MACROS  # noqa: E402

TIEMPO_ORDEN = ["desayuno", "colacion1", "almuerzo", "colacion2", "once", "cena"]
SKIP_SHEETS = {"EDITAR"}
SKIP_FILES = re.compile(r"(?i)^~\$|^copia de|gramajes|planificaci[oó]n cocina|minuta")

# Etiqueta de porción de las planillas → grupo de NutriApp (algunas dependen del tipo de pauta)
GROUP_RULES: List[Tuple[re.Pattern, object]] = [
    (re.compile(r"(?i)^cer"), "cereales"),
    (re.compile(r"(?i)^fruta"), "frutas"),
    (re.compile(r"(?i)^ac?e?ite"), "aceite_grasas"),
    (re.compile(r"(?i)^(g|gr)a?sa"), "alim_ricos_lipidos"),
    (re.compile(r"(?i)^legumbre"), "legumbres"),
    (re.compile(r"(?i)^omega"), "omega_3"),
    (re.compile(r"(?i)^(s\.|sup|suplemento)"), "suplemento_proteico"),
    (re.compile(r"(?i)^l[aá]cteo"), {"vegano": "lacteos_soya", "*": "lacteos_mg"}),
    (re.compile(r"(?i)^(\"?carnes|a\.? ?prot|alimentos? prot|prote[ií]nas)"),
     {"omnivoro": "carnes_bg", "*": "otros_proteicos"}),
]
SIN_PORCION = re.compile(r"(?i)^(verduras?|creatina)")


def parse_qty(text: str) -> Optional[float]:
    """'1', '1/2', '1 1/2', '3', '1,5' → número."""
    try:
        return float(sum(Fraction(p.replace(",", ".")) for p in text.split()))
    except (ValueError, ZeroDivisionError):
        return None


def parse_porcion(label: str, tipo: str) -> Tuple[Optional[str], float, str]:
    """'1 1/2 Legumbre' → ('legumbres', 1.5, '1 1/2 Legumbre')."""
    label = " ".join(str(label).split())
    m = re.match(r"^(\d+(?:[.,]\d+)?(?:\s+\d+/\d+)?|\d+/\d+)\s+(.+)$", label)
    if not m:
        return None, 0.0, label
    qty, name = parse_qty(m.group(1)) or 0.0, m.group(2).strip()
    for pattern, target in GROUP_RULES:
        if pattern.search(name):
            if isinstance(target, dict):
                target = target.get(tipo, target["*"])
            return target, qty, label
    return None, qty, label


def tiempo_de(label: str) -> Optional[str]:
    t = re.sub(r"\s+", "", str(label)).upper()
    numerada = re.match(r"^COMIDA(\d+)$", t)
    if numerada:
        return f"comida{numerada.group(1)}"
    t = re.sub(r"^COMIDA\d+:", "", t)
    for key, word in (("desayuno", "DESAYUNO"), ("almuerzo", "ALMUERZO"), ("snack", "SNACK"),
                      ("snack", "COLACI"), ("cena", "ONCE/CENA"), ("once", "ONCE"), ("cena", "CENA")):
        if t.startswith(word):
            return key
    return None


def titulo_tiempo(label: str) -> str:
    """'D E S A Y U N O (pre…)' → 'Desayuno'; 'COMIDA 2: Snack (trabajo)' → 'Snack'."""
    t = label.split("(")[0].strip()
    if re.fullmatch(r"(\S\s)+\S", t):          # letras espaciadas
        t = t.replace(" ", "")
    t = re.sub(r"(?i)^comida\s*\d+\s*:\s*", "", t).strip()
    return t[:1].upper() + t[1:].lower() if t else "Nota"


def asignar_tiempos(secuencia: List[str]) -> List[str]:
    """Snacks → colación 1 (antes del almuerzo) o colación 2 (después); sin repetir claves."""
    usados, out, visto_almuerzo = set(), [], False
    # "COMIDA 1..N" sin nombre: en orden a desayuno, almuerzo, once, cena
    numeradas = iter(["desayuno", "almuerzo", "once", "cena", "colacion2", "colacion1"])
    secuencia = [next(numeradas) if t.startswith("comida") else t for t in secuencia]
    for t in secuencia:
        if t == "almuerzo":
            visto_almuerzo = True
        key = t if t != "snack" else ("colacion2" if visto_almuerzo else "colacion1")
        if key in usados:
            libres = [k for k in TIEMPO_ORDEN[TIEMPO_ORDEN.index(key):] if k not in usados]
            key = libres[0] if libres else None
        if key:
            usados.add(key)
        out.append(key)
    return out


def _text(v) -> str:
    return " ".join(str(v).split()) if v not in (None, "") else ""


def parse_horizontal(ws, tipo: str):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header = next((i for i, r in enumerate(rows[:6])
                   if any(_text(c).lower().startswith("opci") for c in r)), None)
    if header is None:
        return None
    title_cols = [j for j, c in enumerate(rows[header]) if _text(c).lower().startswith("opci")]
    bloques, actual, notas = [], None, []
    # El primer tiempo de comida puede venir antes de la fila de encabezados
    for r in rows[:header]:
        b = _text(r[1]) if len(r) > 1 else ""
        if b and tiempo_de(b):
            actual = {"tiempo": tiempo_de(b), "porciones": {}, "opciones": [{"plato": "", "lineas": []} for _ in title_cols]}
            bloques = [actual]
    for r in rows[header + 1:]:
        b = _text(r[1]) if len(r) > 1 else ""
        if not b:
            continue
        t = tiempo_de(b)
        grupo, qty, label = parse_porcion(b, tipo)
        if t and not qty:
            extra = re.search(r"\((.+)\)", b)
            if extra:
                notas.append(f"{titulo_tiempo(b)}: {extra.group(1).strip()}")
            actual = {"tiempo": t, "porciones": {}, "opciones": [{"plato": "", "lineas": []} for _ in title_cols]}
            bloques.append(actual)
            continue
        if actual is None:
            continue
        if grupo and qty:
            actual["porciones"][grupo] = actual["porciones"].get(grupo, 0) + qty
        elif not SIN_PORCION.match(label) and not qty:
            continue
        for k, col in enumerate(title_cols):
            plato = _text(r[col]) if len(r) > col else ""
            texto = _text(r[col + 1]) if len(r) > col + 1 else ""
            if plato and not actual["opciones"][k]["plato"]:
                actual["opciones"][k]["plato"] = plato.strip('"').strip()
            if texto and texto != "-":
                actual["opciones"][k]["lineas"].append({"porcion": label, "texto": texto})
    return bloques, notas


def parse_vertical(wb, tipo: str):
    names = {n.strip().upper(): n for n in wb.sheetnames}
    if "PORCIONES" not in names:
        return None
    rows = [list(r) for r in wb[names["PORCIONES"]].iter_rows(values_only=True)]
    header = next((i for i, r in enumerate(rows[:6]) if any("OPCI" in _text(c).upper() for c in r)), None)
    if header is None:
        return None
    opt_cols = [j for j, c in enumerate(rows[header]) if "OPCI" in _text(c).upper()]
    bloques, actual, alt_cols = [], None, {}
    for r in rows[header:]:
        a = _text(r[0]) if r else ""
        t = tiempo_de(a) if a else None
        if t:
            actual = {"tiempo": t, "porciones": {}, "opciones": [{"plato": "", "lineas": []} for _ in opt_cols]}
            bloques.append(actual)
            # Almuerzo con columnas CON/SIN legumbres: cada opción usa la porción de su variante
            alt_cols = {}
            if any("LEGUMBRE" in _text(c).upper() for c in r):
                porc_cols = {_text(r[1]).upper(): 1, _text(r[2]).upper(): 2}
                for j in opt_cols:
                    alt_cols[j] = porc_cols.get(_text(r[j]).upper(), 1)
                continue
        if actual is None or len(r) < 2:
            continue
        label = _text(r[1])
        grupo, qty, label = parse_porcion(label, tipo) if label else (None, 0, "")
        if grupo and qty:
            actual["porciones"][grupo] = actual["porciones"].get(grupo, 0) + qty
        elif not label or not (SIN_PORCION.match(label) or qty):
            continue
        for k, col in enumerate(opt_cols):
            texto = _text(r[col]) if len(r) > col else ""
            porc = _text(r[alt_cols.get(col, 1)]) if alt_cols else label
            if texto:
                actual["opciones"][k]["lineas"].append({"porcion": porc, "texto": texto})
    # Platos desde la hoja MENÚ (en el mismo orden de tiempos)
    menu_sheet = names.get("MENÚ") or names.get("MENU")
    if menu_sheet:
        platos, cur = [], None
        for r in wb[menu_sheet].iter_rows(values_only=True):
            a = _text(r[0]) if r else ""
            if tiempo_de(a):
                cur = [_text(c) for c in r[1:1 + len(opt_cols)]]
                platos.append(cur)
            elif cur is not None and any(_text(c) for c in r[1:1 + len(opt_cols)]):
                for k, c in enumerate(r[1:1 + len(opt_cols)]):
                    if _text(c) and not _text(c).startswith("(*)"):
                        cur[k] = f"{cur[k]} {_text(c)}".strip()
        for bloque, ps in zip(bloques, platos):
            for k, p in enumerate(ps):
                if p and k < len(bloque["opciones"]):
                    bloque["opciones"][k]["plato"] = p
    return bloques, []


def tipo_de(path: str) -> str:
    name = os.path.basename(path).lower()
    if "omn" in path.lower():
        return "omnivoro"
    return "vegano" if name.startswith("vegan") else "ovolacto"


TIPO_LABEL = {"omnivoro": "Omnívoro", "vegano": "Vegano", "ovolacto": "Ovolacto"}


def nombre_de(path: str, sheet: str, multi: bool, tipo: str) -> str:
    base = unicodedata.normalize("NFC", os.path.splitext(os.path.basename(path))[0])
    base = re.sub(r"(?i)^(horizontal\s+)?(vegan|veget)?\s*", "", base)
    base = base.replace(":", "·").replace(" - ", " · ")
    base = " ".join(base.split())
    base = f"{TIPO_LABEL.get(tipo, tipo)} · {base}"
    return f"{base} · {' '.join(sheet.split())}" if multi else base


def to_plantilla(path: str, sheet: str, multi: bool, bloques, notas, tipo: str) -> Optional[dict]:
    numeradas = any(b["tiempo"].startswith("comida") for b in bloques)
    secuencia = asignar_tiempos([b["tiempo"] for b in bloques])
    distribucion, menu = {}, {}
    for key, b in zip(secuencia, bloques):
        if not key:
            continue
        if b["porciones"]:
            distribucion[key] = {g: round(q, 2) for g, q in b["porciones"].items()}
        opciones = [o for o in b["opciones"] if o["lineas"] or o["plato"]]
        if opciones:
            menu[key] = {"opciones": opciones}
    if not distribucion:
        return None
    porciones: Dict[str, float] = {}
    for d in distribucion.values():
        for g, q in d.items():
            porciones[g] = round(porciones.get(g, 0) + q, 2)
    kcal = sum(GRUPOS_MACROS[g]["kcal"] * q for g, q in porciones.items())
    prot = sum(GRUPOS_MACROS[g]["prot"] * q for g, q in porciones.items())
    lip = sum(GRUPOS_MACROS[g]["lip"] * q for g, q in porciones.items())
    m_kcal = re.search(r"(\d{4})\s*k", os.path.basename(path), re.I)
    m_prot = re.search(r"(\d{2,3})\s*g\s*P", os.path.basename(path), re.I)
    resumen = (f"Planilla: {m_kcal.group(1) if m_kcal else '?'} kcal y {m_prot.group(1) if m_prot else '?'} g de proteína. "
               f"Porciones importadas: {kcal:.0f} kcal y {prot:.0f} g de proteína."
               + (" Comidas numeradas: asignadas en orden a desayuno, almuerzo, once y cena." if numeradas else ""))
    return {
        "nombre": nombre_de(path, sheet, multi, tipo)[:255],
        "tipo_pauta": tipo,
        "fa_key": "moderada",
        "ajuste_kcal": None,
        "prot_pct": round(prot * 4 / kcal * 100, 1) if kcal else 20,
        "lip_pct": round(lip * 9 / kcal * 100, 1) if kcal else 30,
        "porciones_json": json.dumps(porciones, ensure_ascii=False),
        "distribucion_json": json.dumps(distribucion, ensure_ascii=False),
        "menu_json": json.dumps(menu, ensure_ascii=False) if menu else None,
        # Las notas se imprimen en el PDF como indicaciones al paciente: solo las de la planilla
        "notes": "\n".join(notas) or None,
        "_origen": f"{path} [{sheet}]",
        "_resumen": resumen,
    }


def leer_carpeta(carpeta: str):
    plantillas, omitidos = [], []
    for path in sorted(glob.glob(os.path.join(carpeta, "**", "*.xlsx"), recursive=True)):
        rel = unicodedata.normalize("NFC", os.path.relpath(path, carpeta))
        if SKIP_FILES.search(os.path.basename(rel)) or "gramajes" in rel.lower():
            omitidos.append((rel, "no es una pauta tipo"))
            continue
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            omitidos.append((rel, f"no se pudo abrir ({type(e).__name__})"))
            continue
        tipo = tipo_de(rel)
        vertical = parse_vertical(wb, tipo)
        if vertical:
            p = to_plantilla(rel, "", False, *vertical, tipo)
            plantillas.append(p) if p else omitidos.append((rel, "sin porciones"))
            continue
        sheets = [s for s in wb.sheetnames if s.strip().upper() not in SKIP_SHEETS]
        found = 0
        for s in sheets:
            res = parse_horizontal(wb[s], tipo)
            if not res:
                continue
            p = to_plantilla(rel, s, len(sheets) > 1, *res, tipo)
            if p:
                plantillas.append(p)
                found += 1
        if not found:
            omitidos.append((rel, "formato no reconocido"))
    # Nombres únicos
    vistos: Dict[str, int] = {}
    for p in plantillas:
        n = p["nombre"]
        vistos[n] = vistos.get(n, 0) + 1
        if vistos[n] > 1:
            p["nombre"] = f"{n} ({vistos[n]})"
    return plantillas, omitidos


def cargar(plantillas, username: str) -> int:
    import models
    from database import SessionLocal
    db = SessionLocal()
    try:
        user = db.query(models.Nutritionist).filter(models.Nutritionist.username == username).one()
        existentes = {p.nombre for p in db.query(models.PlantillaPauta)
                      .filter(models.PlantillaPauta.nutritionist_id == user.id)}
        nuevas = 0
        for p in plantillas:
            if p["nombre"] in existentes:
                continue
            data = {k: v for k, v in p.items() if not k.startswith("_")}
            db.add(models.PlantillaPauta(**data, nutritionist_id=user.id))
            nuevas += 1
        db.commit()
        return nuevas
    finally:
        db.close()


if __name__ == "__main__":
    if len(sys.argv) != 4 or sys.argv[1] not in ("revisar", "cargar"):
        print(__doc__)
        sys.exit(1)
    plantillas, omitidos = leer_carpeta(sys.argv[2])
    print(f"Plantillas: {len(plantillas)}   Archivos omitidos: {len(omitidos)}")
    for rel, motivo in omitidos:
        print(f"  omitido: {rel} → {motivo}")
    if sys.argv[1] == "revisar":
        with open(sys.argv[3], "w", encoding="utf-8") as f:
            json.dump(plantillas, f, ensure_ascii=False, indent=1)
        print(f"Revisión guardada en {sys.argv[3]}")
    else:
        print(f"Cargadas {cargar(plantillas, sys.argv[3])} plantillas nuevas en la cuenta {sys.argv[3]}")
